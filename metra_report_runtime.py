"""Generate standalone Metra Consultation draft reports from Telegram data."""

from __future__ import annotations

import re
import tempfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import inspection_bot as bot


BLUE = RGBColor(31, 78, 121)
LIGHT_BLUE = "D9EAF7"


def _safe_name(value: str) -> str:
    result = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip())
    return result.strip("._") or "metra_report"


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_text(cell, value: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(value or "—"))
    run.bold = bold
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _format_table(table) -> None:
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    table_pr = table._tbl.tblPr
    borders = table_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = tc_pr.first_child_found_in("w:tcMar")
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tc_pr.append(margins)
            for side in ("top", "start", "bottom", "end"):
                margin = margins.find(qn(f"w:{side}"))
                if margin is None:
                    margin = OxmlElement(f"w:{side}")
                    margins.append(margin)
                margin.set(qn("w:w"), "90")
                margin.set(qn("w:type"), "dxa")


def _add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[index], header, bold=True)
        _shade(table.rows[0].cells[index], LIGHT_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            _set_cell_text(cells[index], value)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    _format_table(table)
    return table


def _status_summary(group: dict) -> str:
    statuses = []
    for photo in group.get("photos", []):
        status = photo.get("status", "À confirmer")
        if status not in statuses:
            statuses.append(status)
    return ", ".join(statuses) or "À confirmer"


def _add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def _configure_document(doc: Document, session: dict) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(11)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("METRA CONSULTATION INC.  |  RAPPORT TECHNIQUE")
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.size = Pt(9)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Metra Consultation Inc. — Document de travail à réviser et signer par l’ingénieur responsable")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)


def _add_cover(doc: Document, session: dict, profile: dict) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(55)
    run = paragraph.add_run("METRA")
    run.bold = True
    run.font.size = Pt(34)
    run.font.color.rgb = BLUE
    run = paragraph.add_run("\nCONSULTATION")
    run.bold = True
    run.font.size = Pt(18)

    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(45)
    title = profile["label"].replace(" – ", "\n").replace("–", " ").replace("-", " ")
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0, 0, 0)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(28)
    run = paragraph.add_run(session.get("project_name", "Projet"))
    run.bold = True
    run.font.size = Pt(16)
    paragraph.add_run("\n" + session.get("address", "—"))

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(40)
    paragraph.add_run(f"Date de visite / dossier : {session.get('date', '—')}\n")
    paragraph.add_run("Statut : ÉBAUCHE AUTOMATISÉE — À VALIDER AVANT ÉMISSION").bold = True
    doc.add_page_break()


def _field_rows(session: dict) -> list[list[str]]:
    rows = []
    for value in session.get("project_fields", {}).values():
        if isinstance(value, dict):
            rows.append([value.get("label", "Information"), value.get("value", "—")])
    return rows


def _add_project_definition(doc: Document, session: dict) -> None:
    _add_heading(doc, "1. Identification et définition du projet")
    rows = [
        ["Projet", session.get("project_name", "—")],
        ["Adresse", session.get("address", "—")],
        ["Date", session.get("date", "—")],
    ] + _field_rows(session)
    _add_table(doc, ["Champ", "Information recueillie"], rows, [2.1, 4.8])


def _add_mandate(doc: Document, profile: dict) -> None:
    _add_heading(doc, "2. Mandat et portée")
    doc.add_paragraph(profile["mandate"])
    doc.add_paragraph(
        "Le présent document est une ébauche produite à partir des informations saisies "
        "dans l’outil de terrain. Les documents reçus, les calculs, les hypothèses, les "
        "exigences réglementaires propres au bâtiment et le jugement professionnel doivent "
        "être vérifiés avant toute émission officielle."
    )


def _add_method(doc: Document, type_id: str) -> None:
    _add_heading(doc, "3. Méthodologie et limitations")
    if type_id == "loi16":
        text = (
            "La démarche comprend l’examen des documents disponibles, une visite visuelle "
            "des parties communes accessibles, l’inventaire des composantes, l’évaluation "
            "préliminaire de leur état et de leur durée de vie résiduelle, puis la préparation "
            "d’une planification sur au moins 25 ans. L’étude financière doit ensuite intégrer "
            "les coûts, l’inflation, le rendement du fonds, son solde initial et les dépenses prévues."
        )
    elif type_id == "facade":
        text = (
            "L’inspection vise les composantes accessibles des façades et les conditions "
            "susceptibles d’affecter la sécurité. Les moyens d’accès, sondages, ouvertures "
            "exploratoires et essais effectivement réalisés doivent être décrits dans la version finale."
        )
    elif type_id == "parking":
        text = (
            "L’inspection vise les éléments structuraux et accessoires accessibles du parc de "
            "stationnement, notamment les dalles, poutres, colonnes, murs, joints, drains, "
            "membranes et signes de corrosion, délamination, éclatement ou infiltration."
        )
    else:
        text = (
            "L’inspection porte uniquement sur les équipements identifiés et accessibles. "
            "La version finale doit préciser les essais, charges, instruments, résultats, "
            "exclusions et conditions de remise en service applicables."
        )
    doc.add_paragraph(text)
    doc.add_paragraph(
        "Sauf indication contraire, l’examen est visuel et non destructif. Les parties cachées "
        "ne sont pas réputées inspectées. Les coûts éventuels sont des opinions de planification "
        "et non des soumissions d’entrepreneurs."
    )


def _add_observation_register(doc: Document, session: dict, type_id: str) -> None:
    title = "4. Registre des composantes et observations" if type_id == "loi16" else "4. Sommaire des observations"
    _add_heading(doc, title)
    groups = session.get("groups", [])
    rows = []
    for index, group in enumerate(groups, 1):
        rows.append([
            str(index),
            group.get("element_type", "Élément"),
            _status_summary(group),
            group.get("caption_fr", "Observation à compléter"),
            str(len(group.get("photos", []))),
        ])
    _add_table(
        doc,
        ["No", "Élément / zone", "État", "Observation", "Photos"],
        rows or [["—", "Aucune donnée", "—", "À compléter", "0"]],
        [0.4, 1.55, 1.2, 3.25, 0.55],
    )


def _add_loi16_tables(doc: Document, session: dict) -> None:
    _add_heading(doc, "5. Planification du carnet d’entretien")
    rows = []
    for index, group in enumerate(session.get("groups", []), 1):
        rows.append([
            str(index),
            group.get("element_type", "Composante"),
            _status_summary(group),
            "À confirmer",
            "À confirmer",
            "À confirmer",
        ])
    _add_table(
        doc,
        ["ID", "Composante", "État", "Action", "Échéance", "Coût actuel"],
        rows or [["—", "Inventaire à compléter", "—", "—", "—", "—"]],
        [0.4, 1.7, 1.1, 1.5, 0.9, 1.1],
    )
    doc.add_paragraph(
        "À compléter pour chaque composante : quantité, année d’installation, entretien "
        "préventif, fréquence, durée de vie normale, durée résiduelle, réparation majeure, "
        "année de remplacement, coût, source du coût et documents de référence."
    )

    _add_heading(doc, "6. Étude du fonds de prévoyance")
    fields = session.get("project_fields", {})
    def value(key):
        item = fields.get(key, {})
        return item.get("value", "Non fourni") if isinstance(item, dict) else str(item or "Non fourni")
    _add_table(doc, ["Paramètre", "Valeur / hypothèse"], [
        ["Solde initial du fonds", value("reserve_balance")],
        ["Contribution annuelle actuelle", value("annual_contribution")],
        ["Horizon minimal", "25 ans"],
        ["Inflation des travaux", "À établir et justifier"],
        ["Rendement net du fonds", "À établir et justifier"],
        ["Taxes, honoraires et contingence", "À intégrer selon les hypothèses retenues"],
    ])
    doc.add_paragraph(
        "Équation annuelle à valider : solde de fermeture = solde d’ouverture + contributions "
        "+ rendement net − dépenses planifiées. Un tableau annuel et la contribution recommandée "
        "doivent être joints à la version finale."
    )


def _add_photos(doc: Document, session: dict, heading_number: int) -> None:
    _add_heading(doc, f"{heading_number}. Dossier photographique")
    figure = 1
    for group in session.get("groups", []):
        photos = group.get("photos", [])
        if not photos:
            continue
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(group.get("element_type", "Élément inspecté"))
        run.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        caption = group.get("caption_fr", "Observation à compléter")

        for offset in range(0, len(photos), 2):
            pair = photos[offset:offset + 2]
            table = doc.add_table(rows=1, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for row in table.rows:
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell, photo in zip(table.rows[0].cells, pair):
                path = Path(photo.get("path", ""))
                if not path.exists():
                    _set_cell_text(cell, "[Photo non disponible]")
                    continue
                labelled_path = None
                try:
                    with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as handle:
                        labelled_path = Path(handle.name)
                    bot.add_label_to_image(path, str(figure), labelled_path, display_width_px=900)
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(str(labelled_path), width=Inches(3.15))
                finally:
                    if labelled_path:
                        labelled_path.unlink(missing_ok=True)
                figure += 1
            if len(pair) == 1:
                figure_label = f"Fig. {figure - 1}"
            else:
                figure_label = f"Fig. {figure - len(pair)} à {figure - 1}"
            doc.add_paragraph(f"{figure_label} — {caption}").alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_conclusion(doc: Document, session: dict, heading_number: int) -> None:
    _add_heading(doc, f"{heading_number}. Conclusion préliminaire")
    issues = []
    for group in session.get("groups", []):
        if any(
            "Acceptable" not in photo.get("status", "✅ Acceptable")
            for photo in group.get("photos", [])
        ):
            issue = group.get("caption_fr") or group.get("element_type", "Élément à vérifier")
            issues.append(str(issue).rstrip(" .;"))
    if issues:
        conclusion = doc.add_paragraph(
            "Les observations suivantes requièrent une validation, une intervention ou un suivi "
            "dans la version finale : " + "; ".join(issues) + "."
        )
    else:
        conclusion = doc.add_paragraph(
            "Aucune condition non acceptable n’a été enregistrée dans les données de terrain. "
            "Cette mention ne constitue pas une attestation avant la révision complète du dossier."
        )
    conclusion.paragraph_format.keep_with_next = True
    doc.add_paragraph("Préparé par : Arash Rohani, ing., P.Eng., M.Ing.\nMetra Consultation Inc.")


def build_metra_report(session: dict, lang: str = "fr") -> Path:
    type_id = bot.report_profiles.profile_key(session.get("inspection_type"))
    profile = bot.report_profiles.profile(type_id)
    doc = Document()
    doc.core_properties.author = "Metra Consultation Inc."
    doc.core_properties.title = profile["label"]
    _configure_document(doc, session)
    _add_cover(doc, session, profile)
    _add_project_definition(doc, session)
    _add_mandate(doc, profile)
    _add_method(doc, type_id)
    _add_observation_register(doc, session, type_id)
    if type_id == "loi16":
        _add_loi16_tables(doc, session)
        photo_heading = 7
        conclusion_heading = 8
    else:
        photo_heading = 5
        conclusion_heading = 6
    _add_photos(doc, session, photo_heading)
    _add_conclusion(doc, session, conclusion_heading)

    bot.REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    date = session.get("date", datetime.today().strftime("%Y-%m-%d"))
    output = bot.REPORTS_DIR / f"{_safe_name(session.get('project_name', 'Projet'))}_METRA_{type_id}_{_safe_name(str(date))}_FR.docx"
    doc.save(output)
    return output


def install_metra_report_runtime() -> None:
    original_build_report = bot.build_report

    def build_report(session: dict, lang: str):
        if session.get("company") == "metra":
            return build_metra_report(session, lang)
        return original_build_report(session, lang)

    bot.build_report = build_report
