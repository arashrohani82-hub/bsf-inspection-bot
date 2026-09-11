import os
import sys
import tempfile
import types
import unittest
from pathlib import Path

os.environ.setdefault("TELEGRAM_TOKEN", "test-token")
os.environ.setdefault("ANTHROPIC_API_KEY", "test-key")

if "anthropic" not in sys.modules:
    anthropic_stub = types.ModuleType("anthropic")
    anthropic_stub.Anthropic = lambda **kwargs: object()
    sys.modules["anthropic"] = anthropic_stub

if "telegram" not in sys.modules:
    telegram_stub = types.ModuleType("telegram")
    telegram_stub.Update = type("Update", (), {})
    telegram_stub.ReplyKeyboardMarkup = type("ReplyKeyboardMarkup", (), {})
    telegram_stub.ReplyKeyboardRemove = type("ReplyKeyboardRemove", (), {})
    telegram_ext_stub = types.ModuleType("telegram.ext")
    for name in ("Application", "CommandHandler", "MessageHandler"):
        setattr(telegram_ext_stub, name, type(name, (), {}))
    telegram_ext_stub.ConversationHandler = type("ConversationHandler", (), {"END": -1})
    telegram_ext_stub.ContextTypes = type("ContextTypes", (), {"DEFAULT_TYPE": object})
    telegram_ext_stub.filters = object()
    sys.modules["telegram"] = telegram_stub
    sys.modules["telegram.ext"] = telegram_ext_stub

from docx import Document
from PIL import Image

import inspection_bot as bot
from metra_report_runtime import build_metra_report


class MetraReportTests(unittest.TestCase):
    def test_builds_loi16_draft_with_project_fields_and_photo(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            photo = root / "roof.jpg"
            Image.new("RGB", (900, 600), "lightgray").save(photo)
            bot.REPORTS_DIR = root / "reports"
            session = {
                "company": "metra",
                "inspection_type": "loi16",
                "project_name": "Syndicat Test",
                "address": "123, rue Test, Montréal",
                "date": "2026-09-11",
                "project_fields": {
                    "units": {"label": "Unités", "value": "8"},
                    "reserve_balance": {"label": "Solde du fonds", "value": "25 000 $"},
                },
                "groups": [{
                    "element_type": "Toiture",
                    "caption_fr": "Usure du revêtement à surveiller.",
                    "photos": [{"path": str(photo), "status": "🔧 Réparation requise"}],
                }],
            }
            output = build_metra_report(session)
            self.assertTrue(output.exists())
            document = Document(output)
            full_text = "\n".join(p.text for p in document.paragraphs)
            self.assertIn("METRA", full_text)
            self.assertIn("Étude du fonds de prévoyance", full_text)
            table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            self.assertIn("25 000 $", table_text)
            self.assertIn("Toiture", table_text)


if __name__ == "__main__":
    unittest.main()
